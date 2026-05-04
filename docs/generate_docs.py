import os
from docx import Document
from fpdf import FPDF

content_title = "Multi-Modal AI Content & Spam Detection System"

content_intro = """1. Introduction & Project Overview
The "Multi-Modal AI Content & Spam Detection System" is a unified platform designed to detect malicious or synthetic content across three modalities: Text (Spam), Images (AI-generated/Fake), and Videos (Deepfakes/Manipulated). The system is built with a decoupled architecture, employing separate, highly optimized Machine Learning modules for each data type. The main goals are to provide near real-time predictions, offer a simple and intuitive user interface, and ensure that the backend is robust enough to handle simultaneous multi-modal inference tasks."""

content_arch = """2. System Architecture & Tech Stack
The project adopts a modern web architecture, moving away from simple Flask/HTML to a robust, scalable stack:

- Frontend (UI/UX): Built with Next.js 15 (App Router), TypeScript, and Tailwind CSS. The design philosophy emphasizes a "cybersecurity" and "futuristic" aesthetic using dark themes (#0A0A0A), neon accents (#00D4FF), and glassmorphism. It uses Framer Motion for smooth micro-animations.
- Backend (API & Routing): Powered by FastAPI (Python 3.11). FastAPI was chosen over Flask for its superior asynchronous capabilities, allowing the server to handle concurrent image/video processing without blocking. It includes built-in Swagger documentation.
- Database: MongoDB Atlas (NoSQL) is used with the Motor async driver. It stores user data (via JWT authentication) and logs every inference request for audit trails.
- Machine Learning Layer: 
  - Text: scikit-learn (TF-IDF + Multinomial Naive Bayes)
  - Image: Custom feature extraction (NumPy) + RandomForestClassifier
  - Video: OpenCV for frame extraction, utilizing the Image model."""

content_methodology = """3. Step-by-Step Methodology of Each Module

3.1 Text Spam Detection Module
- Step 1 (Preprocessing): Text is transformed to lowercase, and punctuation/extra whitespace is removed.
- Step 2 (Feature Extraction): The cleaned text is vectorized using TF-IDF (Term Frequency-Inverse Document Frequency) using unigrams and bigrams, up to 10,000 features. This step ensures that statistically significant words (like "winner" or "free") are given more weight than common stop words.
- Step 3 (Classification): The vectorized text is passed to a Multinomial Naive Bayes model.
- Why this approach?: Naive Bayes is extremely fast and effective for text classification. On the SMS Spam dataset, this approach achieved 98.48% accuracy.

3.2 Image Authenticity Detection Module
- Step 1 (Image Loading): The user's uploaded image is resized to 224x224 pixels and converted to a numpy array.
- Step 2 (Statistical Feature Extraction): Instead of using a heavy GPU-dependent Deep Learning model, the system calculates 8 distinct heuristics from the image:
  * Laplacian Variance: Measures blur and natural texture.
  * Colour Entropy: Checks colour histogram diversity.
  * Gradient Magnitude & Local Standard Deviation: Measures natural edge consistency and patch variations.
  * High-Frequency Energy & Minimum Patch Variance: Detects artificially flat regions common in AI-generated images.
  * Colour Diversity & Global Noise Std: Detects the natural noise floor of a real camera.
- Step 3 (Classification): These 8 numerical features are fed into a pre-trained Random Forest Classifier, which outputs a confidence score and a "Real" or "Fake" label.
- Why this approach?: It allows the image detection to run purely on CPU, taking up less than 512MB RAM, making it suitable for free-tier cloud deployment (Render) without sacrificing detection capability.

3.3 Video Deepfake Detection Module
- Step 1 (Frame Extraction): Because video processing is computationally expensive, the system uses OpenCV to extract frames at a rate of 1 frame per second. This dramatically reduces the workload.
- Step 2 (Frame-Level Inference): Each extracted frame is passed independently to the Image Authenticity Module (described above) to generate a "Real" or "Fake" prediction per frame.
- Step 3 (Majority Voting): The system aggregates the predictions of all frames. If more than 50% of the frames are flagged as "Fake", the entire video is flagged as a deepfake. The confidence score is the mean of all frame confidences.
- Why this approach?: It balances accuracy with efficiency. Analyzing every single frame (e.g., 30fps) would crash a basic server; analyzing 1 frame per second provides enough data to spot deepfakes efficiently.

3.4 Backend API & Routing (FastAPI)
- Step 1 (Endpoint Creation): The system defines separate POST endpoints for `/api/detect/text`, `/api/detect/image`, and `/api/detect/video`.
- Step 2 (Asynchronous Processing): FastAPI handles the requests asynchronously. When a video is being processed (which might take 5-10 seconds), the server can still accept text and image requests from other users.
- Step 3 (Logging): After a prediction is made, the result is packaged into a JSON response. Simultaneously, the backend asynchronously writes an audit log to MongoDB containing the timestamp, input type, and result.

3.5 Frontend UI/UX (Next.js)
- Step 1 (User Interaction): The user lands on a dark-themed dashboard. They select their modality (Text, Image, Video) via cards.
- Step 2 (File Handling): A drag-and-drop component handles file uploads, converting images/videos into FormData to be sent to the FastAPI backend.
- Step 3 (State Management): While waiting for the backend, Next.js triggers a Loader component. Once the response is received, a ResultCard component dynamically renders the output (Red for Fake/Spam, Green for Authentic) with the calculated confidence score."""

content_conclusion = """4. Conclusion
By separating the project into distinct, optimized modules, the Multi-Modal AI Detection System achieves high performance without requiring enterprise-level hardware. The use of FastAPI and Next.js ensures a professional, responsive user experience, while the heuristic-based image/video analysis ensures the ML pipeline remains lightweight and deployable."""

def create_docx(filename):
    doc = Document()
    doc.add_heading(content_title, level=0)
    
    doc.add_heading('1. Introduction & Project Overview', level=1)
    doc.add_paragraph(content_intro.split('\n', 1)[1])
    
    doc.add_heading('2. System Architecture & Tech Stack', level=1)
    for line in content_arch.split('\n')[1:]:
        if line.strip():
            if line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('  - '):
                doc.add_paragraph(line[4:], style='List Bullet 2')
            else:
                doc.add_paragraph(line)
                
    doc.add_heading('3. Step-by-Step Methodology of Each Module', level=1)
    
    parts = content_methodology.split('\n\n')
    for part in parts[1:]:
        lines = part.split('\n')
        doc.add_heading(lines[0], level=2)
        for line in lines[1:]:
            if line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('  * '):
                doc.add_paragraph(line[4:], style='List Bullet 2')
            else:
                doc.add_paragraph(line)

    doc.add_heading('4. Conclusion', level=1)
    doc.add_paragraph(content_conclusion.split('\n', 1)[1])
    
    doc.save(filename)
    print(f"Created {filename}")

def create_pdf(filename):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=11)
    
    # Title
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(200, 10, txt=content_title, ln=True, align='C')
    pdf.ln(10)
    
    def add_section(text):
        for line in text.split('\n'):
            if line.strip():
                if line[0].isdigit() and line[1] == '.':
                    pdf.set_font("Arial", 'B', 14)
                    pdf.cell(200, 10, txt=line, ln=True)
                    pdf.set_font("Arial", size=11)
                elif line.startswith('3.') and line[2].isdigit():
                    pdf.set_font("Arial", 'B', 12)
                    pdf.cell(200, 10, txt=line, ln=True)
                    pdf.set_font("Arial", size=11)
                elif line.startswith('- '):
                    pdf.multi_cell(0, 6, txt=f"  - {line[2:]}")
                elif line.startswith('  - ') or line.startswith('  * '):
                    pdf.multi_cell(0, 6, txt=f"      - {line[4:]}")
                else:
                    pdf.multi_cell(0, 6, txt=line)
        pdf.ln(5)

    add_section(content_intro)
    add_section(content_arch)
    add_section(content_methodology)
    add_section(content_conclusion)
    
    pdf.output(filename)
    print(f"Created {filename}")

if __name__ == '__main__':
    create_docx('Project_Methodology_Explanation.docx')
    create_pdf('Project_Methodology_Explanation.pdf')
